"""
ايجنت ذكي للرد التلقائي على استفسارات الزبائن بتليكرام
مبني على Claude API + شيت جوجل (بيانات العروض والأسعار الحية)

طريقة العمل:
1. يستلم رسالة من زبون بتليكرام (عبر Webhook)
2. يقرا آخر نسخة من شيت جوجل (يخزنها مؤقتاً 15 دقيقة، ما يقراها بكل رسالة حتى يكون سريع)
3. يرسل السؤال + بيانات الشيت لـ Claude API
4. يرجع الجواب للزبون بتليكرام
5. لو السؤال مو موجود بالشيت (حجز فعلي، مشكلة، شي غير واضح) → يحول لموظف بشري

الإعداد المطلوب (Environment Variables):
- TELEGRAM_BOT_TOKEN   : توكن البوت من @BotFather
- ANTHROPIC_API_KEY    : مفتاح Claude API
- GOOGLE_SHEET_ID      : أيدي شيت جوجل (من الرابط، الجزء بين /d/ و /edit)
- STAFF_CONTACT        : رقم الموظف اللي يتحول له الزبون (اختياري، فيه قيمة افتراضية)
"""

import os
import time
import io
import requests
import anthropic
from flask import Flask, request
import openpyxl

# ---------------- الإعدادات ----------------
TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY")
GOOGLE_SHEET_ID = os.environ.get("GOOGLE_SHEET_ID")
STAFF_CONTACT = os.environ.get("STAFF_CONTACT", "07711165005 (الكروبات) أو 07810105600 (الطيران)")

if not all([TELEGRAM_BOT_TOKEN, ANTHROPIC_API_KEY, GOOGLE_SHEET_ID]):
    raise RuntimeError(
        "لازم تحط كل من TELEGRAM_BOT_TOKEN و ANTHROPIC_API_KEY و GOOGLE_SHEET_ID "
        "كـ Environment Variables قبل ما تشغل هذا البوت"
    )

TELEGRAM_API = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}"
SHEET_EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{GOOGLE_SHEET_ID}/export?format=xlsx"

claude = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
app = Flask(__name__)

# ---------------- كاش بيانات الشيت (يتحدث كل 15 دقيقة) ----------------
_sheet_cache = {"text": "", "last_fetch": 0}
CACHE_SECONDS = 15 * 60  # 15 دقيقة


def fetch_sheet_text() -> str:
    """يجيب آخر نسخة من شيت جوجل ويحولها لنص مبسط للذكاء الاصطناعي."""
    now = time.time()
    if _sheet_cache["text"] and (now - _sheet_cache["last_fetch"] < CACHE_SECONDS):
        return _sheet_cache["text"]

    try:
        resp = requests.get(SHEET_EXPORT_URL, timeout=30)
        resp.raise_for_status()
        wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)

        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines = [f"### شيت: {sheet_name}"]
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None and str(c).strip() != ""]
                if cells:
                    lines.append(" | ".join(cells))
            # نتجاوز الشيتات الفاضية أو الصغيرة كلش (زوائد)
            if len(lines) > 2:
                parts.append("\n".join(lines))

        text = "\n\n".join(parts)
        # حماية: نقص الطول لو كبر كثير (حتى ما تكبر تكلفة كل رسالة)
        MAX_CHARS = 100_000
        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS] + "\n...[تم اختصار الباقي]"

        _sheet_cache["text"] = text
        _sheet_cache["last_fetch"] = now
        return text

    except Exception as e:
        # لو فشل التحديث، نستخدم النسخة القديمة المخزنة (إذا موجودة) بدل ما نوقف البوت
        if _sheet_cache["text"]:
            return _sheet_cache["text"]
        return f"[تعذر تحميل بيانات العروض حالياً: {e}]"


SYSTEM_PROMPT_TEMPLATE = """أنتَ مساعد خدمة الزبائن لشركة بركات الكوثر للسفر والسياحة، تتكلم بالعراقي بشكل طبيعي وودود ومباشر.

قواعد صارمة يجب اتباعها:
1. جاوب بس من المعلومات الموجودة بجدول العروض تحت. لا تخترع أي سعر أو تاريخ أو معلومة غير موجودة فيه.
2. إذا الزبون سأل عن سعر أو رحلة أو تاريخ مو موجود بالجدول بشكل واضح، لا تخمن السعر — قله بوضوح إنك بتحوله لموظف يتأكد له من السعر الحالي، واعطيه رقم التواصل: {staff_contact}
3. إذا الزبون يريد يأكد حجز فعلي أو يدفع، حوله فوراً لرقم التواصل: {staff_contact} — لا تأكد أي حجز بنفسك.
4. خلي ردودك قصيرة ومباشرة (3-4 جمل كحد أعلى)، بدون مقدمات طويلة.
5. لا تذكر عمولة الشركات أبداً — هذي معلومة داخلية، مو للزبون.
6. إذا الزبون يسلم أو يسأل سؤال عام مالته علاقة بالسفر، جاوبه بلطف وودك.

جدول العروض الحالي:
---
{sheet_data}
---
"""


def ask_claude(user_message: str) -> str:
    sheet_data = fetch_sheet_text()
    system_prompt = SYSTEM_PROMPT_TEMPLATE.format(
        staff_contact=STAFF_CONTACT, sheet_data=sheet_data
    )
    response = claude.messages.create(
        model="claude-sonnet-4-5",
        max_tokens=400,
        system=system_prompt,
        messages=[{"role": "user", "content": user_message}],
    )
    return response.content[0].text


def send_telegram_message(chat_id, text):
    requests.post(
        f"{TELEGRAM_API}/sendMessage",
        json={"chat_id": chat_id, "text": text},
        timeout=15,
    )


# ---------------- Webhook مال تليكرام ----------------
@app.route("/webhook", methods=["POST"])
def telegram_webhook():
    update = request.get_json(silent=True) or {}
    message = update.get("message")
    if not message or "text" not in message:
        return {"ok": True}

    chat_id = message["chat"]["id"]
    user_text = message["text"]

    try:
        reply = ask_claude(user_text)
    except Exception as e:
        reply = f"حصل خطأ فني، تواصل معنا مباشرة: {STAFF_CONTACT}"
        print("ERROR:", e)

    send_telegram_message(chat_id, reply)
    return {"ok": True}


@app.route("/", methods=["GET"])
def health_check():
    return "البوت شغال ✅"


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    app.run(host="0.0.0.0", port=port)
